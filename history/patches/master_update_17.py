"""1.7.0 data update: compare the new Master List (Project Code, Approving Authorities, Codes & Standards) with the live store."""
import json, os, re, sys, glob
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
SP = '/tmp/claude-0/-home-claude/0eb838aa-d16c-5b60-a2e1-bc2abae7b604/scratchpad'
XL = os.path.join(HERE, 'ml-1.7', 'master.xlsx')
CUR = os.path.join(SP, 'live-1.7', 'projects')
OUT = os.path.join(HERE, 'update-1.7'); os.makedirs(OUT, exist_ok=True)

def key_for(h):
    h = re.sub(r'[^a-z0-9]+', ' ', str(h or '').lower()).strip()
    if not h: return None
    rules = [('code', r'^(project|job) ?code'), ('no', r'^(oxypro )?(project ?no|project number|proj ?no|job ?no)'), ('name', r'(project|building).*name'),
             ('sectors', r'sector'), ('levels', r'level|floor'), ('emirate', r'emirate'), ('status', r'status'), ('mainContractor', r'main ?contractor'),
             ('mepConsultant', r'mep ?consultant'), ('leadConsultant', r'lead ?consultant'), ('client', r'^client'), ('scopeType', r'scope ?type'),
             ('scope', r'scope'), ('codes', r'codes?\b.*standard|^codes?$'), ('authorities', r'authorit'), ('value', r'value|amount'), ('duration', r'duration'),
             ('startDate', r'^start'), ('handoverDate', r'handover|complet'), ('year', r'^year'), ('area', r'^area|location|address')]
    for k, rx in rules:
        if re.search(rx, h): return k
    return '?' + h

def clean(k, v):
    if v is None or v == '': return ''
    if hasattr(v, 'year') and k in ('startDate', 'handoverDate'): return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        r = round(v * 100) / 100; return str(int(r)) if r == int(r) else str(r)
    t = str(v)
    if k == 'scope': return '\n'.join(l for l in (re.sub(r'[ \t ]+', ' ', x).strip() for x in re.split(r'\r?\n', t)) if l)
    t = re.sub(r'\s+', ' ', t).strip()
    if k != 'value': t = re.sub(r'\s+,', ',', t); t = re.sub(r',(?=[^\s\d])', ', ', t)
    return t

def norm_status(s):
    s = s or ''
    if re.search(r'complet|handed|finish', s, re.I): return 'Completed'
    if re.search(r'handover|hand over|snag|close ?out', s, re.I): return 'Handover Stage'
    if re.search(r'ongoing|progress|running|current|active|construction', s, re.I): return 'Ongoing'
    return ''

wb = openpyxl.load_workbook(XL, data_only=True); ws = wb['OXYPRO PROJECT MASTER CASE']
hdr = {i: key_for(c.value) for i, c in enumerate(ws[3]) if key_for(c.value)}
print('headers:', hdr)
assert not [v for v in hdr.values() if v.startswith('?')], 'unknown header'
master = {}
for r in range(4, ws.max_row + 1):
    row = [c.value for c in ws[r]]
    no = str(row[0] or '').strip().upper()
    if not re.match(r'^[A-Z0-9][A-Z0-9._-]{0,39}$', no): continue
    d = {k: clean(k, row[i]) for i, k in hdr.items()}
    d['no'] = no; d['emirate'] = d['emirate'].upper(); d['status'] = norm_status(d['status']); d['code'] = d['code'].upper()
    master[no] = d
cur = {json.load(open(f))['no']: json.load(open(f)) for f in glob.glob(os.path.join(CUR, '*.json'))}
ver = {}
for f in glob.glob(os.path.join(CUR, '*.json')):
    pass
print('master', len(master), 'library', len(cur), '| only master', sorted(set(master) - set(cur)), '| only library', sorted(set(cur) - set(master)))
NOW = sys.argv[1] if len(sys.argv) > 1 else '2026-10-27T09:30:00.000Z'
KEYS = [k for k in hdr.values() if k != 'no']
writes, report = [], {}
for no in sorted(master, key=lambda x: int(re.sub(r'\D', '', x) or 0)):
    m, c = master[no], cur[no]
    data = c.get('data', c)
    upd = {}
    for k in KEYS:
        old = str(data.get(k) if data.get(k) is not None else '')
        if m[k] != old and (m[k] or k in data): upd[k] = m[k]
    if 'standards' in data: upd['standards'] = {'__delete__': True}
    report[no] = {k: (str(data.get(k, '')), v) for k, v in upd.items()}
    upd['updatedAt'] = NOW
    p = os.path.join(OUT, f'{no}.json'); json.dump(upd, open(p, 'w'), ensure_ascii=False, indent=1)
    writes.append({'op': 'update', 'collection': 'projects', 'doc_id': no, 'file_path': p, 'if_version': 3})
json.dump(writes, open(os.path.join(OUT, '_batch.json'), 'w'), indent=1)
json.dump(master, open(os.path.join(OUT, '_master.json'), 'w'), ensure_ascii=False, indent=1)
for no, ch in report.items():
    print(f'\n{no}:')
    for k, (o, n) in ch.items(): print(f'  {k}: {o!r} -> {n!r}')
