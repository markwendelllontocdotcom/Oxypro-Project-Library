"""Compare the Oxypro Projects Master List with the library store and write the update batch."""
import json, os, re, sys, glob
import openpyxl

SP = '/tmp/claude-0/-home-claude/0eb838aa-d16c-5b60-a2e1-bc2abae7b604/scratchpad'
XL = os.path.join(SP, 'ml-in', 'master.xlsx')
CUR = os.path.join(SP, 'current', 'projects')
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'update-1.1')
os.makedirs(OUT, exist_ok=True)

COLS = {'A': 'no', 'B': 'name', 'C': 'sectors', 'D': 'levels', 'E': 'area', 'F': 'emirate', 'G': 'mainContractor',
        'H': 'mepConsultant', 'I': 'leadConsultant', 'J': 'client', 'K': 'scope', 'L': 'standards', 'M': 'value',
        'N': 'duration', 'O': 'startDate', 'P': 'handoverDate', 'Q': 'year', 'R': 'status'}
MASTER_KEYS = [k for k in COLS.values() if k != 'no']

def clean(k, v):
    if v is None or v == '':
        return ''
    if hasattr(v, 'year') and k in ('startDate', 'handoverDate'):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        r = round(v * 100) / 100
        return str(int(r)) if r == int(r) else str(r)
    t = str(v)
    if k == 'scope':
        return '\n'.join(l for l in (re.sub(r'[ \t ]+', ' ', x).strip() for x in re.split(r'\r?\n', t)) if l)
    t = re.sub(r'\s+', ' ', t).strip()
    if k != 'value':
        t = re.sub(r'\s+,', ',', t)
        t = re.sub(r',(?=[^\s\d])', ', ', t)
    return t

def norm_status(s):
    s = s or ''
    if re.search(r'complet|handed|finish', s, re.I): return 'Completed'
    if re.search(r'handover|hand over|snag|close ?out', s, re.I): return 'Handover Stage'
    if re.search(r'ongoing|progress|running|current|active|construction', s, re.I): return 'Ongoing'
    return ''

wb = openpyxl.load_workbook(XL, data_only=True)
ws = wb['OXYPRO PROJECT MASTER CASE']
assert [ws[f'{c}3'].value for c in 'AR'] == ['Project No', 'Project Status']
master = {}
for r in range(4, ws.max_row + 1):
    no = ws[f'A{r}'].value
    if not no or not re.match(r'^[A-Z0-9][A-Z0-9._-]{0,39}$', str(no).strip().upper()):
        continue
    d = {k: clean(k, ws[f'{c}{r}'].value) for c, k in COLS.items()}
    d['no'] = str(no).strip().upper()
    d['emirate'] = d['emirate'].upper()
    d['status'] = norm_status(d['status'])
    master[d['no']] = d

cur = {}
for f in glob.glob(os.path.join(CUR, '*.json')):
    d = json.load(open(f)); cur[d['no']] = d

print('master', len(master), 'library', len(cur))
print('only in master:', sorted(set(master) - set(cur)), 'only in library:', sorted(set(cur) - set(master)))
writes, report = [], {}
NOW = sys.argv[1] if len(sys.argv) > 1 else '2026-10-27T05:30:00.000Z'
for no in sorted(master, key=lambda x: int(re.sub(r'\D', '', x) or 0)):
    m = master[no]; c = cur.get(no)
    upd = {}
    for k in MASTER_KEYS:
        old = '' if c is None else str(c.get(k, '') if c.get(k) is not None else '')
        if m[k] != old and (m[k] or k not in (c or {})):
            upd[k] = m[k]
    if c is not None and 'completion' in c:
        upd['completion'] = {'__delete__': True}
    changed = {k: (str((c or {}).get(k, '')), v) for k, v in upd.items() if k != 'completion' and v}
    report[no] = changed
    upd['updatedAt'] = NOW
    path = os.path.join(OUT, f'{no}.json'); json.dump(upd, open(path, 'w'), ensure_ascii=False, indent=1)
    writes.append({'op': 'update', 'collection': 'projects', 'doc_id': no, 'file_path': path, 'if_version': 1})
json.dump(writes, open(os.path.join(OUT, '_batch.json'), 'w'))
json.dump(master, open(os.path.join(OUT, '_master.json'), 'w'), ensure_ascii=False, indent=1)
for no, ch in report.items():
    print(f'\n{no}:')
    for k, (o, n) in ch.items():
        print(f'  {k}: {o!r} -> {n!r}')
